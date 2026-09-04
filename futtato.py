# -*- coding: utf-8 -*-
"""Hajnali Állásvadász — egyesített futtató szkript (v3.0, 2026-09-04).
Ugyanaz a fájl fut a felhőben és a notebookon; az alapmappát az AV_BASE
környezeti változó adja (alapértelmezés: ~/mnt/Claude).

  python3 futtato.py validate <datum>              # nyers_talalatok.json mezőkitöltöttség + javaslat
  python3 futtato.py dedup    <datum>
  python3 futtato.py render   <bemenet.md> <kimenet.docx>   # pandoc docx + ATS formai ellenőrzés
  python3 futtato.py html     <bemenet.md> <kimenet.html>   # Google Doc-nak feltölthető HTML
  python3 futtato.py index    <datum>
  python3 futtato.py tracker  <datum>
  python3 futtato.py manifest <datum>              # feltöltési lista (manifest.json)
  python3 futtato.py xlsx                          # csak a notebookon: csv -> xlsx
  python3 futtato.py sablon                        # ATS docx sablon előállítása (futás elején)

v3.0 változások: a `link` NEM kötelező (pozíció + cég elég); mezőnév-normalizálás
(url/href -> link, company -> ceg, title -> pozicio ...); hírlevél/aggregátor-találatok
kiszűrése; docx helyett HTML->Google Doc feltöltés (manifest).
"""
import sys, os, re, json, csv, zipfile, subprocess, unicodedata, datetime, hashlib
from difflib import SequenceMatcher

VERZIO = "3.0"

BASE    = os.environ.get("AV_BASE") or os.path.join(os.path.expanduser("~"), "mnt", "Claude")
ALAP    = os.path.join(BASE, "00_alap")
NAPI    = os.path.join(BASE, "01_napi")
ALLAPOT = os.path.join(BASE, "02_allapot")
JELENT  = os.path.join(BASE, "03_jelentkezesek")
SEEN    = os.path.join(ALLAPOT, "seen_jobs.json")
NAPLO   = os.path.join(ALLAPOT, "futasi_naplo.md")
REF     = os.path.join(ALAP, "sablon", "ats_reference.docx")
CSV     = os.path.join(JELENT, "jelentkezesek.csv")
SEEN_MAX_NAP = 120           # ennél régebbi dedup-bejegyzést eldobunk

# ---------------------------------------------------------------- segédek
def ma(): return datetime.date.today().isoformat()

def napi_mappa(d=None):
    p = os.path.join(NAPI, d or ma()); os.makedirs(p, exist_ok=True); return p

def ea(s):
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c))

def slug(s, n=45):
    return (re.sub(r"[^a-z0-9]+", "-", ea(s).lower()).strip("-")[:n].strip("-") or "ismeretlen")

def _norm(s): return re.sub(r"[^a-z0-9]+", "", ea(s).lower())

def _poz(s):
    """Pozíció-kulcs: zárójeles rész és hivatkozási számok (KH_104628, #2974068, REF123) nélkül."""
    t = re.sub(r"\([^)]*\)", " ", ea(s or "").lower())
    t = re.sub(r"\b[a-z]{0,4}[_\-#/]?\d{3,}\b", " ", t)
    t = re.sub(r"\b(m/f/d|m/w/d|f/m/x|hu|en)\b", " ", t)
    return re.sub(r"[^a-z0-9]+", "", t)

def _link_kulcs(link):
    """Link-kulcs: séma/www/utm-paraméterek nélkül; kereső-URL (jobs?q=...) nem kulcs."""
    from urllib.parse import urlparse
    try: u = urlparse((link or "").strip())
    except Exception: return ""
    if not u.netloc: return ""
    if re.search(r"(/jobs/?$|/jobs\?|/search|/allashirdetesek/?$)", u.path + "?" + u.query) and "view" not in u.path:
        return ""
    host = u.netloc.lower().replace("www.", "")
    path = u.path.rstrip("/").lower()
    return (host + path) if path else ""

def _ceg(s):
    t = re.sub(r"\b(kft|zrt|nyrt|bt|ltd|gmbh|sa|se|ag|inc|hungary|magyarorszag|hu)\b",
               " ", ea(s).lower())
    return re.sub(r"[^a-z0-9]+", "", t)

def naplo(sor):
    os.makedirs(ALLAPOT, exist_ok=True)
    with open(NAPLO, "a", encoding="utf-8") as f:
        f.write("- `%s` %s\n" % (datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), sor))

def load_profil():
    import yaml
    with open(os.path.join(ALAP, "profil.yaml"), encoding="utf-8") as f:
        return yaml.safe_load(f)

# ---------------------------------------------------------------- 2. fázis
def mar_lattuk(seen, ceg, poz):
    """Fuzzy dedup: hasonló cégnél 88%-os pozíció-egyezés már duplikátum."""
    ck, pk = _ceg(ceg), _poz(poz)
    for k, v in seen.items():
        if not isinstance(v, dict): continue
        vck = v.get("ceg_kulcs"); vpk = v.get("poz_kulcs2") or v.get("poz_kulcs")
        if not (vck and vpk): continue
        if (ck == vck or SequenceMatcher(None, ck, vck).ratio() >= 0.92) \
           and (pk == vpk or SequenceMatcher(None, pk, vpk).ratio() >= 0.88):
            return k
    return None

def mar_lattuk_link(seen, link):
    lk = _link_kulcs(link)
    if not lk: return None
    for k, v in seen.items():
        if isinstance(v, dict) and (v.get("link_kulcs") or _link_kulcs(v.get("link"))) == lk:
            return k
    return None

def prune_seen(seen):
    hatar = (datetime.date.today() - datetime.timedelta(days=SEEN_MAX_NAP)).isoformat()
    return {k: v for k, v in seen.items()
            if not isinstance(v, dict) or (v.get("ujra_latva") or v.get("eloszor", "")) >= hatar}

# Mezőnév-szinonimák: a begyűjtő ágens néha más kulcsot használ. Mindet a kanonikus névre hozzuk.
MEZO_SZINONIMAK = {
    "pozicio": ["pozicio", "pozíció", "position", "title", "cim", "cím", "job_title", "allas", "állás", "megnevezes"],
    "ceg":     ["ceg", "cég", "company", "employer", "munkaltato", "munkáltató", "vallalat", "hirdeto"],
    "helyszin":["helyszin", "helyszín", "location", "hely", "varos", "város", "telephely"],
    "forras":  ["forras", "forrás", "source", "portal", "portál"],
    "link":    ["link", "url", "href", "allas_link", "job_url", "hivatkozas", "hivatkozás", "apply_url"],
    "datum":   ["datum", "dátum", "date", "posted", "kelt"],
    "reszlet": ["reszlet", "részlet", "snippet", "description", "leiras", "leírás", "text", "szoveg", "szöveg"],
    "nyelv":   ["nyelv", "language", "lang"],
}
PORTAL_KULCSOK = {"cvonline", "cvonlinehu", "linkedin", "profession", "professionhu", "indeed",
                  "jooble", "jobline", "nofluffjobs", "glassdoor", "allasportal", "karrierhu"}
URL_RE = re.compile(r"https?://[^\s<>\"')\]]+")

def normalizal(r):
    """Egy nyers találat kulcsait a kanonikus nevekre hozza; a link-et a reszlet-ből is kimenti."""
    if not isinstance(r, dict): return {}
    lower = {(k or "").strip().lower(): v for k, v in r.items()}
    out = {}
    for kanon, nevek in MEZO_SZINONIMAK.items():
        val = ""
        for n in nevek:
            v = lower.get(n.lower())
            if v not in (None, "", [], {}):
                val = v; break
        if isinstance(val, list): val = " ".join(str(x) for x in val if x)
        if isinstance(val, dict): val = val.get("url") or val.get("href") or json.dumps(val, ensure_ascii=False)
        out[kanon] = str(val).strip() if val is not None else ""
    # link: ha üres, de a részletben van URL, azt vesszük
    if not URL_RE.search(out["link"] or ""):
        m = URL_RE.search(out.get("reszlet") or "")
        out["link"] = m.group(0) if m else ""
    # a többi (ismeretlen) mezőt megtartjuk, hogy ne vesszen el információ
    for k, v in r.items():
        kl = (k or "").strip().lower()
        if kl not in {n.lower() for ns in MEZO_SZINONIMAK.values() for n in ns} and k not in out:
            out[k] = v
    return out

def aggregator_e(poz, ceg, link):
    """Hírlevél-összesítő ("15 IT vezető állás Budapest..."), portál-gyökér vagy 'a cég maga a portál'."""
    p = ea(poz).lower()
    if re.match(r"^\s*\d+\s+\S.*\b(allas|allasok|jobs?|ajanlat|talalat)\b", p): return True
    if re.search(r"\b(allasok|allasajanlatok|uj allasok|mentett keresesed|job alerts?|jobs for you)\b", p) \
       and not ceg: return True
    if _ceg(ceg) in PORTAL_KULCSOK and re.search(r"\b(allas|allasok|jobs?)\b", p): return True
    try:
        from urllib.parse import urlparse
        u = urlparse(link or "")
        if u.netloc and u.path.strip("/") in ("", "hu", "en", "allashirdetesek", "jobs", "allas") and not u.query:
            return True
    except Exception: pass
    return False

def _nyers_beolvas(m):
    src = os.path.join(m, "nyers_talalatok.json")
    if not os.path.exists(src): return None, src
    try:
        nyers = json.load(open(src, encoding="utf-8"))
    except Exception as e:
        return {"hiba": "nyers_talalatok.json nem érvényes JSON: %s" % e}, src
    if isinstance(nyers, dict):
        nyers = nyers.get("talalatok") or nyers.get("hirdetesek") or nyers.get("items") or []
    return [normalizal(r) for r in (nyers or []) if isinstance(r, dict)], src

def cmd_validate(datum):
    """A begyűjtő kimenetének gyors egészségügyi ellenőrzése — a vezérlő ez alapján dönt az
    (egyszeri) újra-begyűjtésről, mielőtt a dedup mindent eldobna."""
    m = napi_mappa(datum)
    nyers, src = _nyers_beolvas(m)
    if nyers is None:
        print(json.dumps({"hiba": "nincs nyers_talalatok.json", "utvonal": src, "javaslat": "ujra_begyujtes"},
                         ensure_ascii=False)); return 1
    if isinstance(nyers, dict):
        print(json.dumps({**nyers, "javaslat": "ujra_begyujtes"}, ensure_ascii=False)); return 1
    n = len(nyers)
    def ar(k): return round(100.0 * sum(1 for r in nyers if (r.get(k) or "").strip()) / n) if n else 0
    stat = {"bejovo": n, "pozicio_pct": ar("pozicio"), "ceg_pct": ar("ceg"), "link_pct": ar("link"),
            "helyszin_pct": ar("helyszin"), "reszlet_pct": ar("reszlet"),
            "aggregator": sum(1 for r in nyers if aggregator_e(r.get("pozicio",""), r.get("ceg",""), r.get("link",""))),
            "forrasok": {}}
    for r in nyers:
        f = r.get("forras") or "ismeretlen"; stat["forrasok"][f] = stat["forrasok"].get(f, 0) + 1
    if n == 0: stat["javaslat"] = "ures"
    elif stat["pozicio_pct"] < 60 or (stat["ceg_pct"] < 40 and stat["link_pct"] < 40):
        stat["javaslat"] = "ujra_begyujtes"
        stat["indok"] = "a találatok többségénél hiányzik a pozíció, vagy a cég ÉS a link is"
    elif stat["link_pct"] < 30:
        stat["javaslat"] = "ok_link_nelkul"
        stat["indok"] = "kevés a link — a dedup átengedi, a csomagoknál a levélbeli részletből dolgozunk"
    else:
        stat["javaslat"] = "ok"
    # normalizált változat mentése, hogy a dedup és az ember is ugyanazt lássa
    json.dump({"datum": datum, "talalatok": nyers}, open(os.path.join(m, "nyers_normalizalt.json"), "w",
              encoding="utf-8"), ensure_ascii=False, indent=1)
    naplo("1b. fázis (validate): %d találat, pozíció %d%%, cég %d%%, link %d%%, aggregátor %d → %s"
          % (n, stat["pozicio_pct"], stat["ceg_pct"], stat["link_pct"], stat["aggregator"], stat["javaslat"]))
    print(json.dumps(stat, ensure_ascii=False)); return 0

def cmd_dedup(datum):
    m = napi_mappa(datum)
    nyers, src = _nyers_beolvas(m)
    if nyers is None:
        print(json.dumps({"hiba": "nincs nyers_talalatok.json", "utvonal": src}, ensure_ascii=False)); return 1
    if isinstance(nyers, dict):
        print(json.dumps(nyers, ensure_ascii=False)); return 1

    P = load_profil()
    kizaro = [ea(k).lower() for k in (P.get("kizaro_kulcsszavak") or [])]
    lok = P.get("lokacio", {}) or {}
    elso = [ea(x).lower() for x in lok.get("elsodleges", [])]
    masod = (lok.get("masodlagos_csak_hibrid_vagy_ho") or {})
    masod_ter = [ea(x).lower() for x in masod.get("teruletek", [])]
    tav_szavak = [ea(w) for w in ["home office","homeoffice","távmunka","remote","hibrid",
                                  "hybrid","otthonról","distance","anywhere"]]

    try:
        seen = json.load(open(SEEN, encoding="utf-8")) if os.path.exists(SEEN) else {}
        if not isinstance(seen, dict): seen = {}
    except Exception:
        seen = {}
    ki, st = [], {"bejovo": len(nyers), "duplikatum": 0, "kizaro_szo": 0,
                  "lokacio": 0, "hianyos": 0, "aggregator": 0, "link_nelkul": 0, "atmegy": 0}

    for r in nyers:
        poz = (r.get("pozicio") or "").strip(); ceg = (r.get("ceg") or "").strip()
        link = (r.get("link") or "").strip()
        # v3: a link NEM kötelező. Hiányos = nincs pozíció, VAGY nincs cég ÉS nincs link sem.
        if not poz or (not ceg and not link): st["hianyos"] += 1; continue
        if aggregator_e(poz, ceg, link): st["aggregator"] += 1; continue
        if not link:
            r["link"] = ""; r["link_hianyzik"] = True; st["link_nelkul"] += 1
        if any(k in ea("%s %s" % (poz, r.get("reszlet",""))).lower() for k in kizaro):
            st["kizaro_szo"] += 1; continue

        hely = ea(r.get("helyszin") or "").lower()
        tav = any(w in ea("%s %s" % (hely, r.get("reszlet",""))).lower() for w in tav_szavak)
        ok = (not hely) or any(x in hely for x in elso) or tav
        if any(x in hely for x in masod_ter) and not tav: ok = False
        if not ok: st["lokacio"] += 1; continue

        kulcs = _norm("%s|%s" % (ceg, poz))
        talalt = kulcs if kulcs in seen else (mar_lattuk_link(seen, link) or mar_lattuk(seen, ceg, poz))
        if talalt:
            seen[talalt]["ujra_latva"] = datum; st["duplikatum"] += 1
            # ha a korábbi példánynak nem volt linkje, a mostanit megjegyezzük
            if link and not seen[talalt].get("link"):
                seen[talalt]["link"] = link; seen[talalt]["link_kulcs"] = _link_kulcs(link)
            continue
        seen[kulcs] = {"pozicio": poz, "ceg": ceg, "eloszor": datum, "link": link,
                       "ceg_kulcs": _ceg(ceg), "poz_kulcs": _norm(poz), "poz_kulcs2": _poz(poz),
                       "link_kulcs": _link_kulcs(link)}
        r["dedup_kulcs"] = kulcs; ki.append(r); st["atmegy"] += 1

    os.makedirs(ALLAPOT, exist_ok=True)
    json.dump(prune_seen(seen), open(SEEN, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    out = os.path.join(m, "szurt_talalatok.json")
    json.dump({"datum": datum, "statisztika": st, "talalatok": ki},
              open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    naplo("2. fázis (dedup): %d be → %d tovább (dupla %d, kizáró szó %d, lokáció %d, hiányos %d, aggregátor %d, link nélkül %d)"
          % (st["bejovo"], st["atmegy"], st["duplikatum"], st["kizaro_szo"], st["lokacio"],
             st["hianyos"], st["aggregator"], st["link_nelkul"]))
    print(json.dumps({"kimenet": out, **st}, ensure_ascii=False)); return 0

# ---------------------------------------------------------------- 5. fázis
def cmd_render(md_path, docx_path=None):
    docx_path = docx_path or os.path.splitext(md_path)[0] + ".docx"
    cmd = ["pandoc", md_path, "-f", "markdown+smart", "-t", "docx", "-o", docx_path, "--wrap=none"]
    if os.path.exists(REF): cmd += ["--reference-doc", REF]
    subprocess.run(cmd, check=True)
    r = ellenoriz(docx_path)
    print(json.dumps(r, ensure_ascii=False, indent=1))
    return 2 if r["hibak"] else 0

def cmd_html(md_path, html_path=None):
    """Markdown -> egyszerű HTML, amit a Drive `textContent`-tel Google Doc-ká konvertál.
    Ez váltja ki a base64-es docx-feltöltést (az volt a csomagok elvesztésének oka)."""
    html_path = html_path or os.path.splitext(md_path)[0] + ".html"
    body = subprocess.run(["pandoc", md_path, "-f", "markdown+smart", "-t", "html", "--wrap=none"],
                          check=True, capture_output=True, text=True).stdout
    cim = os.path.splitext(os.path.basename(md_path))[0]
    doc = ('<!DOCTYPE html><html lang="hu"><head><meta charset="utf-8"><title>%s</title>'
           '<style>body{font-family:Calibri,Arial,sans-serif;font-size:10.5pt}'
           'h1{font-size:17pt}h2{font-size:13pt}h3{font-size:11.5pt}</style></head><body>%s</body></html>'
           % (cim, body))
    open(html_path, "w", encoding="utf-8").write(doc)
    print(json.dumps({"fajl": html_path, "bajt": os.path.getsize(html_path)}, ensure_ascii=False)); return 0

def _ceg_fajlnev(mappa_nev):
    """'2_karrier-hungaria_head-of-it' -> 'Karrier_Hungaria' (a docx/Doc névhez)."""
    resz = mappa_nev.split("_", 2)
    ceg = resz[1] if len(resz) > 1 else mappa_nev
    return "_".join(w.capitalize() for w in ceg.split("-") if w) or "Ceg"

def cmd_manifest(datum):
    """Pontos feltöltési lista a 7. fázisnak. A feltöltő ágens ezt követi, és CSAK azt tölti fel,
    ami a Drive napi mappájában még nincs (név szerint) — így nincs duplikátum és nem marad ki semmi."""
    m = napi_mappa(datum)
    tetelek = []
    def add(rel, mime, mod, cim=None, gyoker="napi"):
        p = os.path.join(m, rel) if gyoker == "napi" else rel
        if os.path.exists(p):
            tetelek.append({"helyi": p, "mappa": os.path.dirname(rel) if gyoker == "napi" else gyoker,
                            "cim": cim or os.path.basename(rel), "mime": mime, "mod": mod,
                            "bajt": os.path.getsize(p)})
    # napi gyökér — nyers fájlok is, hogy hibánál visszanézhető legyen
    add("index.html", "text/html", "fajl")
    add("talalatok.json", "application/json", "fajl")
    add("nyers_talalatok.json", "application/json", "fajl")
    add("nyers_normalizalt.json", "application/json", "fajl")
    add("szurt_talalatok.json", "application/json", "fajl")
    add("pontozott.json", "application/json", "fajl")
    add("tovabbi_talalatok.md", "text/markdown", "fajl")
    if os.path.exists(os.path.join(m, "osszefoglalo.md")):
        cmd_html(os.path.join(m, "osszefoglalo.md"), os.path.join(m, "osszefoglalo.html"))
        add("osszefoglalo.html", "text/html", "gdoc", cim="osszefoglalo")
    # csomagmappák
    for d in sorted(os.listdir(m)):
        dp = os.path.join(m, d)
        if not os.path.isdir(dp) or not re.match(r"^\d+_", d): continue
        for fn in ("hirdetes.md", "cv.md", "level.md", "beadas.md", "ats_riport.md"):
            add(os.path.join(d, fn), "text/markdown", "fajl")
        ceg = _ceg_fajlnev(d)
        for src, cim in (("cv.md", "Hornyak_Sandor_CV_%s" % ceg),
                         ("level.md", "Hornyak_Sandor_motivacios_level_%s" % ceg)):
            sp = os.path.join(dp, src)
            if os.path.exists(sp):
                hp = os.path.join(dp, os.path.splitext(src)[0] + ".html")
                cmd_html(sp, hp)
                add(os.path.join(d, os.path.basename(hp)), "text/html", "gdoc", cim=cim)
    # állapotfájlok (dátumozott név, mert a Drive nem tud felülírni)
    add(SEEN, "application/json", "fajl", cim="seen_jobs_%s.json" % datum, gyoker="02_allapot")
    add(CSV, "text/csv", "fajl", cim="jelentkezesek_%s.csv" % datum, gyoker="03_jelentkezesek")
    out = os.path.join(m, "manifest.json")
    json.dump({"datum": datum, "verzio": VERZIO, "tetelek": tetelek}, open(out, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    osszes = sum(t["bajt"] for t in tetelek)
    naplo("7. fázis előkészítés: manifest %d tétel, %d bájt." % (len(tetelek), osszes))
    print(json.dumps({"manifest": out, "tetel": len(tetelek), "bajt": osszes,
                      "gdoc": [t["cim"] for t in tetelek if t["mod"] == "gdoc"]}, ensure_ascii=False)); return 0

def ellenoriz(docx_path):
    hibak, figy, info = [], [], {}
    with zipfile.ZipFile(docx_path) as z:
        names = z.namelist(); doc = z.read("word/document.xml").decode("utf-8", "ignore")
        if "<w:tbl>" in doc: hibak.append("Táblázat a dokumentumban — az ATS parser összekeveri a sorrendet.")
        if re.search(r'<w:cols[^>]*w:num="([2-9])"', doc): hibak.append("Többhasábos elrendezés.")
        if any(n.startswith("word/media/") for n in names): hibak.append("Kép/grafika a dokumentumban.")
        if "<w:txbxContent>" in doc or "<v:textbox" in doc: hibak.append("Szövegdoboz a dokumentumban.")
        for part in [n for n in names if re.match(r"word/(header|footer)\d*\.xml", n)]:
            t = re.sub(r"<[^>]+>", "", z.read(part).decode("utf-8", "ignore")).strip()
            if t: hibak.append("Nem üres élőfej/élőláb (%s): '%s'" % (part, t[:60]))
        sz = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", doc)).strip()
        info["karakter"] = len(sz); info["becsult_oldal"] = round(len(sz) / 3300.0, 1)
        if info["becsult_oldal"] > 2.6:
            figy.append("Becsült terjedelem %s oldal — vezetői CV-nél 2 az ideális." % info["becsult_oldal"])
        if "@" not in sz: hibak.append("Nincs e-mail cím a törzsben.")
        if not re.search(r"\+?\s*36\s*\)?[\s\-/]*\d", sz): figy.append("Nem találtam telefonszámot a törzsben.")
        regi = sorted(set(re.findall(r"\b(19[5-8]\d)\b", sz)))
        if regi: hibak.append("Régi évszám(ok): %s — életkorra utalnak, ki kell venni." % ", ".join(regi))
        if re.search(r'w:color w:val="FFFFFF"', doc):
            hibak.append("Fehér (rejtett) szöveg — az ATS-ek kiszűrik és kizárnak.")
    return {"fajl": os.path.basename(docx_path), "rendben": not hibak,
            "hibak": hibak, "figyelmeztetesek": figy, **info}

# ---------------------------------------------------------------- 6. fázis
def cmd_index(datum):
    import html
    m = napi_mappa(datum); src = os.path.join(m, "talalatok.json")
    if not os.path.exists(src): print("nincs talalatok.json"); return 1
    D = json.load(open(src, encoding="utf-8"))
    T = sorted(D.get("talalatok", []), key=lambda x: -(x.get("pontszam") or 0))
    e = lambda s: html.escape(str(s or ""))
    def sav(p):
        return ("teljes","Teljes csomag") if p>=85 else ("reszleges","Levél + summary") \
               if p>=65 else ("lista","Csak lista") if p>=40 else ("gyenge","Alacsony")
    sorok=[]
    for t in T:
        p=t.get("pontszam") or 0; kl,cimke=sav(p); mp=t.get("mappa")
        fajlok=""
        if mp and os.path.isdir(os.path.join(m,mp)):
            fajlok='<div class="files">'+" · ".join(
                '<a href="%s/%s">%s</a>'%(e(mp),e(fn),e(fn))
                for fn in sorted(os.listdir(os.path.join(m,mp))))+'</div>'
        ats=t.get("ats") or {}; lf=ats.get("lefedettseg")
        atsb='<span class="pill %s">ATS %s%%</span>'%("ok" if (lf or 0)>=75 else "warn",e(lf)) if ats else ""
        b=t.get("beadas") or {}
        bl=('<a class="apply" href="%s" target="_blank" rel="noopener">Beadás: %s</a>'
            %(e(b["url"]),e(b.get("csatorna") or "link"))) if b.get("url") else \
           ('<span class="apply muted">Beadás: %s</span>'%e(b["csatorna"]) if b.get("csatorna") else "")
        sorok.append("""
<article class="job %s"><div class="score"><b>%s</b><span>%s</span></div><div class="main">
<h3><a href="%s" target="_blank" rel="noopener">%s</a></h3>
<div class="meta">%s · %s · <span class="src">%s</span>%s</div>
<p class="why">%s</p>%s<div class="row">%s%s</div>%s</div></article>""" % (
            kl,e(p),e(cimke),e(t.get("link")),e(t.get("pozicio")),e(t.get("ceg")),
            e(t.get("helyszin")),e(t.get("forras")),
            " · "+e(t.get("ber")) if t.get("ber") else "", e(t.get("indoklas")),
            '<p class="gap"><b>Hiányzik:</b> %s</p>'%e(t.get("hiany")) if t.get("hiany") else "",
            atsb,bl,fajlok))
    st=D.get("statisztika") or {}
    statsor=" · ".join("%s: %s"%(k,v) for k,v in st.items())
    kim=D.get("kimaradt_forrasok") or []
    doc = INDEX_SABLON % {"datum":e(datum),"stat":e(statsor),
        "torzs":"".join(sorok) if sorok else '<div class="empty">Ma nem érkezett új, releváns hirdetés.</div>',
        "kimaradt":'<p><b>Kimaradt források:</b> %s</p>'%e(", ".join(kim)) if kim else "",
        "ido":datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}
    out=os.path.join(m,"index.html"); open(out,"w",encoding="utf-8").write(doc)
    naplo("6. fázis: index.html kész, %d találat." % len(T)); print("OK ->",out); return 0

INDEX_SABLON = """<!doctype html><html lang="hu"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Állások — %(datum)s</title><style>
:root{--bg:#f4f4f1;--sf:#fbfbf9;--ink:#191d1b;--ink2:#4a5150;--ink3:#767d7b;
--rule:#d8d9d4;--ac:#0d6f66;--warn:#a2541b;--good:#3d6b2e}
@media(prefers-color-scheme:dark){:root{--bg:#131715;--sf:#1a1f1d;--ink:#e7eae7;--ink2:#b0b8b5;
--ink3:#848d8a;--rule:#2e3532;--ac:#43b3a4;--warn:#d99a55;--good:#8cc177}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:15px/1.55 -apple-system,Segoe UI,Roboto,sans-serif}
.wrap{max-width:900px;margin:0 auto;padding:28px 20px 70px}
h1{font-size:26px;margin:0 0 4px;letter-spacing:-.02em}
.sub{color:var(--ink3);font-size:13px;margin:0 0 22px;font-family:ui-monospace,Menlo,monospace}
.job{display:grid;grid-template-columns:78px 1fr;gap:16px;background:var(--sf);
border:1px solid var(--rule);border-radius:8px;padding:14px 16px;margin-bottom:12px}
.job.teljes{border-left:4px solid var(--ac)}.job.reszleges{border-left:4px solid var(--warn)}
.job.lista,.job.gyenge{border-left:4px solid var(--rule)}
.score{text-align:center;padding-top:2px}
.score b{display:block;font-size:26px;line-height:1;font-variant-numeric:tabular-nums}
.score span{font-size:10px;color:var(--ink3);display:block;margin-top:4px}
h3{margin:0 0 3px;font-size:17px;line-height:1.25}
h3 a{color:var(--ink);text-decoration:none}h3 a:hover{color:var(--ac)}
.meta{font-size:12.5px;color:var(--ink2);margin-bottom:7px}
.src{font-family:ui-monospace,Menlo,monospace;font-size:11.5px;color:var(--ink3)}
.why{margin:0 0 6px;color:var(--ink2);font-size:14px}
.gap{margin:0 0 6px;font-size:13px;color:var(--warn)}
.row{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:6px}
.pill{font-size:11px;font-family:ui-monospace,monospace;padding:2px 7px;border-radius:3px;border:1px solid var(--rule)}
.pill.ok{color:var(--good);border-color:var(--good)}.pill.warn{color:var(--warn);border-color:var(--warn)}
.apply{font-size:12.5px;color:var(--ac)}.apply.muted{color:var(--ink3)}
.files{font-size:11.5px;font-family:ui-monospace,Menlo,monospace;color:var(--ink3);
padding-top:6px;border-top:1px dashed var(--rule)}
.files a{color:var(--ac);text-decoration:none}.files a:hover{text-decoration:underline}
footer{margin-top:26px;padding-top:16px;border-top:1px solid var(--rule);color:var(--ink3);font-size:12.5px}
.empty{background:var(--sf);border:1px solid var(--rule);border-radius:8px;padding:26px;text-align:center;color:var(--ink3)}
</style></head><body><div class="wrap">
<h1>Állások — %(datum)s</h1><p class="sub">%(stat)s</p>%(torzs)s
<footer>%(kimaradt)s<p>Generálva: %(ido)s · A rendszer semmit nem adott be a nevedben.</p></footer>
</div></body></html>"""

# ---------------------------------------------------------------- tracker
FEJ = ["Azonosító","Első nap","Pontszám","Pozíció","Cég","Helyszín","Forrás","Hirdetés linkje",
       "Beadás csatornája","Beadás linkje","ATS %","Csomag mappa","Státusz","Beadva",
       "Visszajelzés","Megjegyzés"]

def cmd_tracker(datum):
    m = napi_mappa(datum); src = os.path.join(m, "talalatok.json")
    if not os.path.exists(src): print("nincs talalatok.json"); return 1
    T = [t for t in json.load(open(src, encoding="utf-8")).get("talalatok", []) if t.get("mappa")]
    os.makedirs(JELENT, exist_ok=True)
    sorok, meglevo = [], set()
    if os.path.exists(CSV):
        with open(CSV, encoding="utf-8", newline="") as f:
            rd = list(csv.reader(f))
        if rd:
            sorok = rd[1:]; meglevo = {r[0] for r in sorok if r}
    uj = 0
    for t in T:
        az = t.get("dedup_kulcs") or "%s|%s" % (t.get("ceg"), t.get("pozicio"))
        if az in meglevo: continue
        b = t.get("beadas") or {}
        sorok.append([az, datum, t.get("pontszam"), t.get("pozicio"), t.get("ceg"),
                      t.get("helyszin"), t.get("forras"), t.get("link"), b.get("csatorna"),
                      b.get("url"), (t.get("ats") or {}).get("lefedettseg"), t.get("mappa"),
                      "Előkészítve", "", "", ""]); uj += 1
    with open(CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(FEJ); w.writerows(sorok)
    naplo("Tracker: %d új sor a jelentkezesek.csv-be." % uj)
    print(json.dumps({"fajl": CSV, "uj_sor": uj}, ensure_ascii=False)); return 0

def cmd_xlsx():
    """Csak a notebookon: a csv-ből formázott xlsx-et készít."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if not os.path.exists(CSV): print("nincs jelentkezesek.csv"); return 1
    rows = list(csv.reader(open(CSV, encoding="utf-8", newline="")))
    wb = Workbook(); ws = wb.active; ws.title = "Jelentkezések"
    for r in rows: ws.append(r)
    for i, w in enumerate([16,11,9,34,24,18,13,42,20,42,7,30,14,11,18,30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D6F66")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    out = os.path.join(JELENT, "jelentkezesek.xlsx"); wb.save(out)
    print("OK ->", out); return 0


# ---------------------------------------------------------------- sablon
def cmd_sablon():
    """ATS-barát pandoc reference.docx előállítása a semmiből.
    Determinisztikus, ezért nem kell tárolni: minden környezet előállítja magának."""
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    tmp = "/tmp/_pandoc_ref_default.docx"
    subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                   stdout=open(tmp, "wb"), check=True)
    d = Document(tmp)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                     s.even_page_header, s.even_page_footer):
            try:
                for pa in part.paragraphs: pa.text = ""
            except Exception: pass
    BLACK = RGBColor(0, 0, 0)
    def st(name, size, bold=False, italic=False, sb=None, sa=None, kwn=False):
        try: style = d.styles[name]
        except KeyError: return
        f = style.font
        f.name = "Calibri"; f.size = Pt(size); f.bold = bold; f.italic = italic
        f.color.rgb = BLACK; f.all_caps = False
        pf = style.paragraph_format
        if sb is not None: pf.space_before = Pt(sb)
        if sa is not None: pf.space_after = Pt(sa)
        pf.keep_with_next = kwn; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for n in ("Normal", "Body Text", "First Paragraph"): st(n, 10.5, sb=0, sa=4)
    st("Compact", 10.5, sb=0, sa=2)
    st("Title", 17, bold=True, sb=0, sa=2)
    st("Subtitle", 11, sb=0, sa=8)
    st("Heading 1", 13, bold=True, sb=12, sa=4, kwn=True)
    st("Heading 2", 11.5, bold=True, sb=9, sa=3, kwn=True)
    st("Heading 3", 10.5, bold=True, italic=True, sb=7, sa=2, kwn=True)
    for n in ("Author", "Date"): st(n, 10.5, sb=0, sa=2)
    os.makedirs(os.path.dirname(REF), exist_ok=True)
    d.save(REF); print("OK ->", REF); return 0

# ---------------------------------------------------------------- belépés
if __name__ == "__main__":
    a = sys.argv[1:] 
    if not a: print(__doc__); sys.exit(1)
    c = a[0]
    try:
        if   c == "validate": sys.exit(cmd_validate(a[1] if len(a) > 1 else ma()))
        elif c == "dedup":    sys.exit(cmd_dedup(a[1] if len(a) > 1 else ma()))
        elif c == "render":   sys.exit(cmd_render(a[1], a[2] if len(a) > 2 else None))
        elif c == "html":     sys.exit(cmd_html(a[1], a[2] if len(a) > 2 else None))
        elif c == "index":    sys.exit(cmd_index(a[1] if len(a) > 1 else ma()))
        elif c == "tracker":  sys.exit(cmd_tracker(a[1] if len(a) > 1 else ma()))
        elif c == "manifest": sys.exit(cmd_manifest(a[1] if len(a) > 1 else ma()))
        elif c == "xlsx":     sys.exit(cmd_xlsx())
        elif c == "sablon":   sys.exit(cmd_sablon())
        elif c == "verzio":   print(VERZIO); sys.exit(0)
        else: print("ismeretlen parancs:", c); sys.exit(1)
    except SystemExit: raise
    except Exception as e:
        # Sose haljon meg csendben: a vezérlő JSON-ként lássa a hibát, és naplózva legyen.
        naplo("HIBA (%s): %s" % (c, e))
        print(json.dumps({"hiba": str(e), "parancs": c}, ensure_ascii=False)); sys.exit(3)
